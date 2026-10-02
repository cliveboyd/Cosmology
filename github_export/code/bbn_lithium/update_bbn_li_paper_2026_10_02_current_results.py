from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
SOURCE = (
    ROOT
    / "github_export"
    / "docs"
    / "papers"
    / "bbn_lithium_selective_electromagnetic_transfer_candidate_2026-09-25_release_candidate_v1.docx"
)
OUTPUT = (
    ROOT
    / "github_export"
    / "docs"
    / "papers"
    / "bbn_lithium_selective_electromagnetic_transfer_candidate_2026-10-02_release_candidate_v2.docx"
)


def find_paragraph(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_paragraph_start(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph prefix not found: {prefix}")


def paragraph_index(doc: Document, target) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph._p is target._p:
            return index
    raise ValueError(f"Paragraph is not in document: {target.text}")


def replace_text(paragraph, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph.clear()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(colour)
    run_properties.append(underline)
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is not None:
        return
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def format_table(table, widths: list[float], body_size: float = 7.4) -> None:
    table.autofit = False
    repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = 1
            if row_index == 0:
                shade_cell(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(body_size if row_index else body_size + 0.2)
                    if row_index == 0:
                        run.bold = True


def append_table_row(table, values: list[str], body_size: float = 7.2) -> None:
    cells = table.add_row().cells
    for cell, value in zip(cells, values):
        cell.text = value
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(body_size)
    prevent_row_split(table.rows[-1])


def make_paragraph(doc: Document, text: str, style: str = "Normal"):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    paragraph.paragraph_format.keep_together = True
    return paragraph


def make_heading(doc: Document, text: str, level: int):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def move_before(reference_paragraph, elements) -> None:
    for element in elements:
        reference_paragraph._p.addprevious(element)


def add_october_section(doc: Document) -> None:
    discussion = find_paragraph(doc, "Discussion")
    blocks = []

    heading = make_heading(doc, "October 2026 Action-Derived Source and Cascade Qualification", 1)
    heading.paragraph_format.page_break_before = True
    blocks.append(heading._p)
    blocks.append(
        make_paragraph(
            doc,
            "The 28 September to 2 October programme replaced the earlier response-level source with progressively more physical tests. The sequence was deliberately ordered: fluence calibration, direct LINX discretisation checks, a minimal action-derived lifetime and shut-off law, exact ACROPOLIS decay cascades, production and cosmology screens, and independent reduced-cascade reproduction. Each pass and failure is retained because they constrain different parts of the proposed mechanism.",
        )._p
    )

    blocks.append(make_heading(doc, "Full-Fluence and Direct-Network Qualification", 2)._p)
    blocks.append(
        make_paragraph(
            doc,
            "The fluence-matched NLO-proxy PRIMAT run completed 10,000/10,000 draws. It retained 1,785 precision-nine passes and 1,433 rows that passed every registered D/H branch. These fractions describe a hypothesis-search distribution under the declared partial covariance; they are not a posterior probability for the mediator model.",
        )._p
    )
    blocks.append(
        make_paragraph(
            doc,
            "A stratified direct LINX rerun then completed four convergence-controlled rows, including the source-off control. Three rows passed the technical validation criterion, the largest abundance difference was 0.325 sigma, and two of the three source rows retained the same precision-nine and all-D/H classifications. The low-temperature-tail audit subsequently showed that the best row was numerically converged but source-prescription sensitive: 82.0% of its baseline exposure came from the low-temperature tail, and neither floor removal nor cutoff robustness passed. This converted an apparent numerical success into a physical-source question.",
        )._p
    )

    blocks.append(make_heading(doc, "Action, Production and Exact ACROPOLIS", 2)._p)
    blocks.append(
        make_paragraph(
            doc,
            "The minimal scalar action used V(phi) = m_phi^2 phi^2 / 2 and L_phi-gamma = -g_phi-gamma phi F_mn F^mn / 4. Expansion dilution and the lifetime were derived from this ansatz, while the nuclear response normalisation remained calibrated to the frozen LINX endpoint. None of the 126 primary s-wave, p-wave or mild inverse-speed rows passed; only diagnostic responses rising approximately as T^beta with beta no larger than -3 produced passes. The minimal EFT source therefore failed and a resonance, threshold enhancement or additional production channel became necessary.",
        )._p
    )
    blocks.append(
        make_paragraph(
            doc,
            "The native ACROPOLIS 1.3.1 scan evaluated 15,768 decay-mediator rows without a runtime failure. Fifty-eight rows passed the precision-nine and all-D/H gates. The best conditional row had m_phi = 4.4 MeV, E_gamma = 2.2 MeV, tau = 150,000 s, n_phi/n_gamma = 1.53993 x 10^-6, D/H = 2.512784 x 10^-5, Yp = 0.2469982, Li7/H = 1.238279 x 10^-10 and abundance chi2 = 0.904474. This is an exact cascade and photodisintegration result for a scanned initial scalar abundance, not a derivation of that abundance.",
        )._p
    )
    blocks.append(
        make_paragraph(
            doc,
            "When the 58 rows were replayed against 10,000 correlated modern PRIMAT abundance draws, every row retained at least one abundance pass; the largest precision-nine and all-D/H pass fractions were 27.95% and 11.40%. The eta, N_eff and FIRAS visibility screens were largely benign, but inverse-decay freeze-in supplied at most 45.3% of the required scalar population. No row passed the combined production screen. A separately declared parent-decay source then produced a candidate-17 bridge that passed all eight registered staging gates, with Li7/H = 1.291358 x 10^-10 and chi2 = 1.305824, but its reheating branching ratio was calibrated rather than independently predicted.",
        )._p
    )

    caption = make_paragraph(doc, "Table 12. September 28 to October 2 qualification chronology")
    caption.runs[0].bold = True
    caption.paragraph_format.keep_with_next = True
    blocks.append(caption._p)
    table = doc.add_table(rows=1, cols=4)
    table.style = doc.tables[15].style
    headers = ["Stage", "Quantitative result", "Gate status", "Scientific meaning"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    rows = [
        ["10,000-draw full-fluence NLO proxy", "1,785 precision-nine; 1,433 all-D/H passes", "Hypothesis-search complete", "Large candidate basin under the declared proxy; not a model posterior."],
        ["Stratified direct LINX", "4/4 completed; 3 technical passes; max 0.325 sigma", "Technical partial pass", "Discretisation is controlled for the tested rows, but classifications are not universal."],
        ["Low-temperature tail", "82.0% baseline tail exposure; cutoff and floor-removal gates fail", "Gate incomplete", "The source history, rather than solver convergence, controls the candidate."],
        ["Minimal action-derived source", "0/126 primary rows pass", "Failed", "Ordinary dilution plus minimal response does not sustain the required late source."],
        ["Exact ACROPOLIS scan", "15,768 rows; 58 all-D/H passes; best chi2 0.9045", "Conditional parameter pass", "A viable exact decay-cascade basin exists when the initial scalar abundance is scanned."],
        ["Modern abundance and cosmology replay", "Max all-D/H pass 11.4%; freeze-in/required <= 0.453", "Production gate failed", "Abundance and visibility can pass, but the minimal action underproduces the source."],
        ["Parent-decay action bridge", "8/8 staging gates; Li7/H 1.2914 x 10^-10", "Technical pass", "Time staging closes conditionally; the parent branching remains calibrated."],
        ["Independent reduced basin", "40/58 classifications agree; 54/58 within 20%; 11-row robust core", "Global gate incomplete", "Candidate 17 is robust, but hard pair-opening treatment fails across the full basin."],
        ["Explicit pair-production audit", "11/11 core; 6/11 neighbours; all 22 within 10%", "Targeted technical pass", "Angle-integrated Breit-Wheeler treatment removes the hard-cutoff artefact for the core."],
        ["Nuclear-active-band redistribution", "5/5 boundary rows recovered; 4/5 controls retained; max 1.588%", "Technical gate incomplete", "Photon-photon redistribution explains most residual discrepancy; one boundary remains."],
        ["basin_052 boundary audit", "GALAH two-sided tail 0.0481; ESS interval spans the 2.5% boundary", "Boundary unresolved", "The lone control flip is too close to observational and cross-section boundaries for a binary claim."],
    ]
    for values in rows:
        append_table_row(table, values)
    format_table(table, [1.12, 1.78, 1.12, 2.48])
    blocks.append(table._tbl)

    blocks.append(make_heading(doc, "Independent Cascade and Boundary Readout", 2)._p)
    blocks.append(
        make_paragraph(
            doc,
            "The independent reduced line cascade reproduced the action-history candidate at Li7/H = 1.287252 x 10^-10, only 0.318% from the ACROPOLIS bridge, with stable cutoff and resolution checks. Across the full 58-row ACROPOLIS pass basin, however, only 40 classifications agreed and 54 remained within 20% in lithium; 11 rows formed a diagnostic robust core. This failure was traced mainly to the hard pair-opening convention rather than numerical resolution.",
        )._p
    )
    blocks.append(
        make_paragraph(
            doc,
            "Replacing that convention with the angle-integrated Breit-Wheeler rate retained all 11 robust-core classifications and 6/11 adjacent neighbours, with all 22 rows within 9.97% of ACROPOLIS. Adding repeated photon-photon redistribution in the Be7-active band recovered all five remaining boundary rows and brought all ten targeted rows within 1.588% of ACROPOLIS, but retained only four of five paired controls. The technical gate therefore remains incomplete rather than failed wholesale.",
        )._p
    )
    blocks.append(
        make_paragraph(
            doc,
            "For the sole flipped control, basin_052, the primary lithium value lies at GALAH empirical CDF 0.02403, with a two-sided tail probability of 0.04805. The autocorrelation-ESS interval for that CDF is 0.01711-0.03371 and therefore crosses the registered 0.025 lower boundary. A -0.104% change in the Be7 photodisintegration scale reaches the boundary. The preferred BRICK/de Souza summary proxy places 46.27% of 500,000 draws inside the GALAH interval, reproduced at 46.20% with an independent seed. The earlier binary failure is consequently model-sensitive and unresolved, not a confirmed pass or robust rejection.",
        )._p
    )

    caption = make_paragraph(doc, "Table 13. Current promotion gates after the October cascade programme")
    caption.runs[0].bold = True
    caption.paragraph_format.keep_with_next = True
    blocks.append(caption._p)
    table = doc.add_table(rows=1, cols=4)
    table.style = doc.tables[11].style
    headers = ["Gate", "Current evidence", "Status", "Required closure"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    rows = [
        ["Abundance target", "GALAH posterior plus 10,000-draw PRIMAT replay", "Qualified observational target", "Independent stellar-survey and depletion replication."],
        ["Ordinary D-chain uncertainty", "Best earlier candidate D_M^2 = 136.914; 0/60 in partial 95% region", "Failed under registered mapping", "Exact synchronised LGLL/TUNL posterior products."],
        ["Exact mediator cascade", "58 ACROPOLIS all-D/H passes", "Conditional parameter pass", "Derive rather than scan the mediator population."],
        ["Source action", "Minimal inverse decay underproduces; calibrated parent bridge passes", "Incomplete", "Predict parent production, branching and phase-space history from one frozen action."],
        ["Independent cascade", "Candidate 17 agrees at 0.318%; active-band boundary agreement <= 1.588%", "Targeted technical pass", "External full coupled electron-photon cascade and independent nuclear response."],
        ["Mass-7 uncertainty", "basin_052 boundary spans GALAH ESS interval; BRICK/de Souza used only as summary proxy", "Unresolved", "Independent Be7 photodisintegration covariance and exact synchronised mass-7 rate draws."],
        ["CMB and spectral safety", "Planck qualified; eta/N_eff and FIRAS visibility screens benign for exact pass rows", "Screen pass, likelihood open", "Full FIRAS covariance and ACT/SPT-consistent likelihood closure."],
    ]
    for values in rows:
        append_table_row(table, values)
    format_table(table, [1.25, 2.35, 1.1, 1.8])
    blocks.append(table._tbl)

    tv = make_paragraph(
        doc,
        "TV index: October evidence updates correspond to the full-fluence, action-lifetime, ACROPOLIS, production, parent-bridge, independent-cascade, explicit-pair, active-band and basin_052 registered outputs dated 28 September to 2 October 2026.",
    )
    tv.runs[0].italic = True
    blocks.append(tv._p)

    move_before(discussion, blocks)


def update_cover_and_narrative(doc: Document) -> None:
    cover = doc.paragraphs
    replace_text(cover[1], "October 2026 action-derived source, exact ACROPOLIS and independent-cascade update")
    cover[1].runs[0].font.size = Pt(11)
    replace_text(cover[2], "Planck PR4/NPIPE, GALAH, external covariance and cascade qualification retained")
    cover[2].runs[0].font.size = Pt(11)
    replace_text(cover[4], "Generated: 2 October 2026 | Melbourne, Australia")
    cover[4].runs[0].font.size = Pt(11)
    replace_text(
        cover[6],
        "EVIDENCE STATUS: Formal paper update. The Planck, GALAH, 10,000-draw fluence, exact ACROPOLIS, action-history bridge and targeted independent-cascade calculations are technically recorded.",
        bold_prefix="EVIDENCE STATUS:",
    )
    replace_text(
        cover[7],
        "SCIENTIFIC STATUS: Gate incomplete. The exact cascade admits conditional abundance-compatible parameter space, but the minimal action underproduces the mediator population, the full independent-cascade basin does not close, and ordinary D-chain uncertainty remains strongly disfavoured.",
        bold_prefix="SCIENTIFIC STATUS:",
    )
    replace_text(
        cover[8],
        "PROMOTION BOUNDARY: no lithium solution or discovery is claimed. Promotion requires a frozen production action, exact external D-chain and mass-7 posterior products, an external coupled cascade or independent nuclear response, and full FIRAS/ACT/SPT consistency.",
        bold_prefix="PROMOTION BOUNDARY:",
    )

    replace_text(find_paragraph(doc, "Latest Results - Overview"), "Latest Results - Overview through 2 October 2026")
    replace_text(
        find_paragraph_start(doc, "The newer stellar analysis makes the lithium target narrower."),
        "The GALAH DR3/DR4 analysis retains a censored stellar-surface target of A(Li) = 2.09062, corresponding to Li/H = 1.232 x 10^-10, with a 95% interval of 1.073-1.424 x 10^-10. It remains an empirical stellar anchor rather than a direct primordial-abundance measurement.",
    )
    replace_text(
        find_paragraph_start(doc, "A large computer search then found many mathematical combinations"),
        "The later programme moved beyond the September rate-pull search. A 10,000-draw fluence-matched PRIMAT run, convergence-controlled LINX propagation and an exact 15,768-row ACROPOLIS decay scan established conditional abundance-compatible parameter space. Fifty-eight ACROPOLIS rows passed every registered D/H branch; the best had Li7/H = 1.2383 x 10^-10 and abundance chi2 = 0.9045.",
    )
    replace_text(
        find_paragraph_start(doc, "The current result therefore separates two explanations."),
        "The physical interpretation is narrower. The minimal scalar action did not generate enough late source, while a calibrated parent-decay bridge and the best independent reduced cascade passed technically. The full 58-row independent basin did not close, although explicit pair production and active-band photon redistribution reduced the selected discrepancies to below 1.588%. Ordinary nuclear-rate error remains disfavoured, and the remaining mass-7 boundary is unresolved pending external posterior products.",
    )

    abstract_texts = [
        "A staged investigation of the cosmological lithium problem is reported in a reference-integrated form.",
        "July 2026 broad BBN and cosmological controls produced no joint D/H, helium and Li-7 solution; August identified a selective post-BBN electromagnetic-transfer candidate.",
        "September Planck PR4/NPIPE, CMB-blackbody and GALAH analyses retained a constrained observational target, while an external D-chain covariance audit rejected ordinary quoted rate uncertainty as the required explanation under the registered mapping.",
        "The present update completes a 10,000-draw fluence-matched NLO-proxy run, convergence-controlled direct LINX checks, an action-conditioned lifetime audit, an exact ACROPOLIS scan, modern-abundance and cosmology replay, a parent-decay source bridge and independent cascade tests.",
        "The exact ACROPOLIS scan evaluated 15,768 rows and found 58 all-D/H passes. The best conditional row has m_phi = 4.4 MeV, tau = 150,000 s, Li7/H = 1.238279 x 10^-10 and abundance chi2 = 0.904474.",
        "The minimal inverse-decay action supplies at most 45.3% of the required scalar abundance, so it fails the production gate. A calibrated parent-decay bridge passes its staging gates and gives Li7/H = 1.291358 x 10^-10, but its branching ratio is not independently predicted.",
        "An independent reduced cascade reproduces the selected action-history candidate to 0.318%, yet only 40/58 classifications agree over the full basin. Explicit Breit-Wheeler and active-band redistribution calculations recover the targeted boundaries to within 1.588% of ACROPOLIS, with one paired control unresolved.",
        "That control, basin_052, lies at GALAH empirical CDF 0.02403; its ESS-aware interval spans the preregistered 0.025 boundary and a 0.104% Be7 cross-section shift changes the classification.",
        "The paper therefore reports a conditional exact-cascade parameter basin and a progressively narrowed source mechanism, not a solution. Scientific promotion remains gated by a predictive production action, exact external D-chain and mass-7 posterior products, an external full coupled cascade or independent nuclear response, and complete FIRAS/ACT/SPT closure.",
    ]
    abstract_heading = find_paragraph(doc, "Abstract")
    index = paragraph_index(doc, abstract_heading)
    for offset, text in enumerate(abstract_texts, 1):
        replace_text(doc.paragraphs[index + offset], text)

    replace_text(
        find_paragraph_start(doc, "This paper records the transition from negative July controls"),
        "This paper records the transition from negative July controls, through the positive August candidate and September CMB, stellar and covariance qualification, to the 28 September-2 October source-action and cascade programme. The chronology matters: each later calculation tests a previously open physical link and preserves both successful and failed gates.",
    )
    replace_text(
        find_paragraph_start(doc, "The computational programme used LINX"),
        "The computational programme used LINX and PRIMAT for BBN network and abundance validation, ACROPOLIS 1.3.1 for exact electromagnetic decay cascades, Cobaya and CAMB for Planck likelihood sampling, and project-specific source-history, production, transport and covariance audits. The October sequence also implemented independent reduced Compton-recycling, angle-integrated Breit-Wheeler and photon-photon redistribution calculations to localise disagreements with ACROPOLIS.",
    )
    replace_text(
        find_paragraph_start(doc, "The external literature determines how these modules are interpreted"),
        "External references determine the interpretation: LINX and PRIMAT define network controls; ACROPOLIS defines the full-cascade benchmark; Planck and FIRAS define cosmological and spectral gates; GALAH defines the current stellar-surface lithium likelihood; and Pisanti/LUNA, LGLL, TUNL, BRICK and de Souza define the independent nuclear-uncertainty products required for promotion.",
    )

    discussion = find_paragraph(doc, "Discussion")
    i = paragraph_index(doc, discussion)
    discussion_texts = [
        "The October calculations materially strengthen the candidate while also making its remaining weakness more precise. An exact ACROPOLIS decay-cascade basin exists, the best abundance row is quantitatively strong, and the action-history candidate survives a targeted independent transport calculation.",
        "The result is not explained by ordinary quoted D-chain uncertainty, and it is not generated by the minimal inverse-decay action. The model now requires an explicit production sector or parent population whose branching, abundance and momentum history are predicted rather than calibrated.",
        "The independent cascade programme shows why a single validation row is insufficient. Candidate 17 is stable, but the wider basin is sensitive to pair-production and photon-redistribution modelling near the Be7 threshold. Adding these processes recovers most targeted discrepancies and leaves one observationally marginal control rather than a global cascade failure.",
        "The basin_052 audit is especially important for claim discipline. Its classification changes under sub-percent Be7 cross-section movement and its GALAH ESS interval straddles the registered lower boundary. It should therefore be reported as unresolved, not counted as either decisive support or decisive falsification.",
        "The current hypothesis is consequently testable and narrower than the August proposal: a late electromagnetic source in the Be7-open, deuterium-closed window can work conditionally, but its population, full cascade and nuclear posterior must be closed by independent inputs.",
        "This distinction preserves the central scientific result. The programme has progressed from a curve-fitting candidate to an exact-cascade parameter basin with identified production and boundary requirements, while stopping short of a solution claim.",
    ]
    for offset, text in enumerate(discussion_texts, 1):
        replace_text(doc.paragraphs[i + offset], text)

    limitations = find_paragraph(doc, "Limitations")
    i = paragraph_index(doc, limitations)
    limitations_texts = [
        "The exact ACROPOLIS scan conditions on a scanned initial scalar abundance; the minimal action does not derive enough of that abundance.",
        "The parent-decay bridge uses a calibrated reheating branching ratio and a cohort approximation rather than a fully sampled particle-production model.",
        "The independent cascade is reduced rather than a second full coupled electron-photon Boltzmann package, and it reuses the published Be7 response fit.",
        "The basin_052 mass-7 calculation uses independent Gaussian summary blocks because exact synchronised BRICK and de Souza rate curves and an independent Be7 photodisintegration covariance are not available locally.",
        "The public FIRAS table is used with visibility screens because the full correlated covariance is not yet registered, and ACT/SPT likelihood closure remains open.",
        "The Planck analysis is spectrum-level rather than map-level, and the abundance and CMB gates are not yet sampled in one global microscopic posterior.",
        "The GALAH DR4 check remains a same-survey release validation; independent stellar-depletion and external survey replication are pending.",
        "All search-enriched pass fractions are descriptive of the tested samples and must not be interpreted as discovery significances or model probabilities.",
    ]
    for offset, text in enumerate(limitations_texts, 1):
        replace_text(doc.paragraphs[i + offset], text)

    conclusions = find_paragraph(doc, "Conclusions")
    i = paragraph_index(doc, conclusions)
    conclusion_texts = [
        "The July investigations failed to resolve the lithium problem with broad BBN or cosmological modifications.",
        "The August selective electromagnetic candidate survived the September Planck, blackbody and GALAH gates, while the external D-chain covariance audit rejected ordinary quoted rate uncertainty as its explanation.",
        "The October work establishes a conditional exact ACROPOLIS basin: 58 of 15,768 rows pass every registered D/H branch, with a best abundance chi2 of 0.904474.",
        "The mechanism is not yet closed. Minimal inverse-decay production is insufficient; the calibrated parent bridge is technically successful but not predictive; the full independent cascade basin is incomplete; and basin_052 remains an observational and mass-7 cross-section boundary case.",
        "The strongest defensible conclusion is therefore that a narrow late electromagnetic source remains a quantitatively viable conditional candidate with identified falsification gates. It is not yet a solution of the cosmological lithium problem.",
    ]
    for offset, text in enumerate(conclusion_texts, 1):
        replace_text(doc.paragraphs[i + offset], text)

    future = find_paragraph(doc, "Future Work")
    i = paragraph_index(doc, future)
    future_texts = [
        "Freeze one parent-production action and predict the reheating yield, branching ratio, mediator phase-space distribution and late two-photon source without calibration to the abundance target.",
        "Reproduce the candidate and boundary rows in an external full coupled electron-photon cascade or an independently implemented nuclear-response pipeline.",
        "Obtain exact synchronised LGLL/TUNL D-chain posterior products and independent Be7 photodisintegration covariance or posterior samples; then repeat the joint gate.",
        "Propagate exact BRICK and de Souza mass-7 posterior curves rather than independent Gaussian summary blocks.",
        "Run full FIRAS covariance, ACT DR6 and SPT-3G likelihood checks with parameter and calibration consistency.",
        "Replicate the GALAH lithium posterior in an independent stellar survey with a physical depletion model.",
        "Publish a frozen reproducibility bundle with checksums for the action declaration, source history, exact cascade, independent transport, abundance inputs and final paper outputs.",
    ]
    for offset, text in enumerate(future_texts, 1):
        replace_text(doc.paragraphs[i + offset], text)

    priority = find_paragraph(doc, "Highest-priority next tests")
    i = paragraph_index(doc, priority)
    priority_texts = [
        "Obtain an external Be7 photodisintegration covariance or posterior sample set and exact synchronised BRICK/de Souza rate curves, then repeat basin_052 without the summary-response approximation.",
        "Implement the frozen parent-decay source in a second full cascade or independent nuclear-response calculation and preregister the candidate and boundary rows.",
        "Replace the calibrated parent branching with a predictive production calculation from the same action, including momentum history, inverse processes, entropy and energy accounting.",
        "Obtain synchronised LGLL Gaussian-process or paired TUNL D-chain samples on a common temperature grid and repeat the external covariance gate.",
        "Run official FIRAS covariance plus ACT DR6/SPT-3G likelihood closure while retaining the Planck, D/H, Yp, He3/D and GALAH gates unchanged.",
        "Retain all failed and boundary rows in the public manifest so future source models are tested against the same frozen evidence set.",
    ]
    for offset, text in enumerate(priority_texts, 1):
        replace_text(doc.paragraphs[i + offset], text)


def update_gate_and_manifest_tables(doc: Document) -> None:
    methods = doc.tables[2]
    methods.rows[4].cells[1].text = (
        "Cobaya/CAMB with Planck PR4/NPIPE CamSpec, low-l TT, low-l EE and PR4 lensing; "
        "split-chain, autocorrelation-ESS and registered-cell density audit."
    )
    methods.rows[4].cells[2].text = (
        "Replaces the compressed CMB gate and determines whether the native likelihood "
        "is technically qualified."
    )
    methods._tbl.remove(methods.rows[5]._tr)
    format_table(methods, [1.65, 2.3, 2.55], body_size=7.0)

    gate_table = doc.tables[11]
    gate_table.rows[2].cells[2].text = (
        "Conditionally satisfied in 58 exact ACROPOLIS rows and in the action-history candidate; modern PRIMAT replay gives up to 11.4% all-D/H pass."
    )
    gate_table.rows[2].cells[3].text = (
        "Repeat with a predictive source and exact synchronised external D-chain posterior draws."
    )
    gate_table.rows[4].cells[2].text = (
        "Exact ACROPOLIS basin found; candidate 17 independently reproduced to 0.318%; full reduced basin and one active-band control remain open."
    )
    gate_table.rows[4].cells[3].text = (
        "External full coupled cascade or independent nuclear response, then full FIRAS covariance."
    )
    format_table(gate_table, [1.15, 1.75, 2.0, 1.6], body_size=7.1)

    programs = doc.tables[22]
    additions = [
        ["run_bbn_li_primat_nlo_fullfluence.py", "Runs the 10,000-draw full-fluence-matched NLO-proxy robustness branch."],
        ["diagnose_bbn_li_fullfluence_linx_validation.py", "Performs stratified direct LINX source validation with segmented convergence control."],
        ["diagnose_bbn_li_action_lifetime_shutoff.py", "Derives the conditional scalar lifetime/dilution history and tests source shut-off responses."],
        ["run_bbn_li_acropolis_mediator.py", "Runs the exact ACROPOLIS 1.3.1 non-universal decay-cascade parameter scan."],
        ["diagnose_bbn_li_acropolis_cosmology_closure.py", "Replays exact pass rows against modern PRIMAT, production, eta/N_eff and FIRAS screens."],
        ["diagnose_bbn_li_parent_decay_acropolis_bridge.py", "Bridges the action-conditioned parent source into native ACROPOLIS evolution."],
        ["diagnose_bbn_li_independent_line_cascade.py", "Implements the independent reduced line-cascade candidate check."],
        ["run_bbn_li_independent_cascade_basin_overnight.py", "Extends reduced-cascade comparison across all 58 exact ACROPOLIS pass rows."],
        ["diagnose_bbn_li_explicit_pair_cascade.py", "Adds angle-integrated Breit-Wheeler pair production and secondary-energy closure."],
        ["diagnose_bbn_li_active_band_cascade.py", "Solves the Be7-active photon-band redistribution recurrence for boundary rows."],
        ["diagnose_bbn_li_basin052_boundary_robustness.py", "Audits GALAH ESS and mass-7 cross-section sensitivity for the remaining control."],
    ]
    existing = {row.cells[0].text for row in programs.rows}
    for values in additions:
        if values[0] not in existing:
            append_table_row(programs, values, body_size=7.0)
    format_table(programs, [2.65, 3.85], body_size=7.0)

    objects = doc.tables[23]
    object_additions = [
        ["Exact ACROPOLIS mediator scan", "plamb_runs/diagnostics/bbn_li_acropolis_mediator/"],
        ["Action and production closure", "plamb_runs/diagnostics/bbn_li_action_lifetime_shutoff/; bbn_li_parent_decay_acropolis_bridge/"],
        ["Independent cascade sequence", "plamb_runs/diagnostics/bbn_li_independent_line_cascade/; bbn_li_independent_cascade_basin/; bbn_li_explicit_pair_cascade/"],
        ["Active-band and boundary audit", "plamb_runs/diagnostics/bbn_li_active_band_cascade/; bbn_li_basin052_boundary_robustness/"],
        ["GitHub October result bundle", "github_export/results/2026-10-02/bbn_lithium/"],
    ]
    existing = {row.cells[0].text for row in objects.rows}
    for values in object_additions:
        if values[0] not in existing:
            append_table_row(objects, values, body_size=7.0)
    format_table(objects, [2.45, 4.05], body_size=7.0)


def add_references(doc: Document) -> None:
    anchor = find_paragraph(doc, "Additional September 2026 Anchor References")
    entries = [
        (
            "D. Odell et al., Performing Bayesian analyses with AZURE2 using BRICK: an application to the 7Be system, Frontiers in Physics 10, 888476 (2022). ",
            "https://arxiv.org/abs/2112.12838",
        ),
        (
            "R. S. de Souza et al., Hierarchical Bayesian Thermonuclear Rate for the 7Be(n,p)7Li Big Bang Nucleosynthesis Reaction, ApJ 894, 134 (2020). ",
            "https://arxiv.org/abs/1912.06210",
        ),
    ]
    for text, url in entries:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.add_run(text)
        add_hyperlink(paragraph, "[link]", url)
        anchor._p.addprevious(paragraph._p)


def clean_empty_heading_before_discussion(doc: Document) -> None:
    discussion = find_paragraph(doc, "Discussion")
    current = discussion._p.getprevious()
    while current is not None and current.tag == qn("w:p"):
        texts = current.xpath(".//w:t/text()")
        if "".join(texts).strip():
            break
        previous = current.getprevious()
        parent = current.getparent()
        parent.remove(current)
        current = previous


def clean_empty_paragraphs_before_results(doc: Document) -> None:
    results = find_paragraph(doc, "Results")
    current = results._p.getprevious()
    while current is not None and current.tag == qn("w:p"):
        texts = current.xpath(".//w:t/text()")
        if "".join(texts).strip():
            break
        previous = current.getprevious()
        parent = current.getparent()
        parent.remove(current)
        current = previous


def apply_accessibility_metadata(doc: Document) -> None:
    for table in doc.tables:
        if table.rows:
            repeat_table_header(table.rows[0])

    for index, drawing_properties in enumerate(
        doc.element.xpath(".//*[local-name()='docPr']"), start=1
    ):
        if not drawing_properties.get("descr"):
            drawing_properties.set(
                "descr",
                "BBN-Li posterior qualification and diagnostic comparison charts.",
            )
        if not drawing_properties.get("title"):
            drawing_properties.set("title", f"BBN-Li diagnostic figure {index}")


def ensure_page_numbers_are_retained(doc: Document) -> None:
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                if "Updated September 25, " in run.text:
                    run.text = run.text.replace(
                        "Updated September 25, ", "Updated October 2, "
                    )
        footer_text = " ".join(p.text for p in section.footer.paragraphs)
        footer_xml = section.footer._element.xml
        if "PAGE" not in footer_xml and not footer_text.strip():
            paragraph = section.footer.paragraphs[0]
            paragraph.alignment = 2
            paragraph.add_run("Page ")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            paragraph._p.append(field)


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    update_cover_and_narrative(doc)
    update_gate_and_manifest_tables(doc)
    add_references(doc)
    add_october_section(doc)
    clean_empty_paragraphs_before_results(doc)
    clean_empty_heading_before_discussion(doc)
    apply_accessibility_metadata(doc)
    ensure_page_numbers_are_retained(doc)

    properties = doc.core_properties
    properties.title = "A Selective Electromagnetic-Transfer Candidate for the Cosmological Lithium Problem"
    properties.subject = "October 2026 current-results update: exact ACROPOLIS and independent-cascade qualification"
    properties.author = "Clive Stewart Boyd; David Ng"
    properties.keywords = "BBN, lithium-7, beryllium-7, ACROPOLIS, LINX, PRIMAT, electromagnetic cascade"
    properties.comments = (
        "Release candidate v2. Conditional exact-cascade parameter basin; scientific promotion remains gated."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Saved DOCX: {OUTPUT}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
